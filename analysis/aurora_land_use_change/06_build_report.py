#!/usr/bin/env python3
"""Create the illustrated Aurora data-center land-use change report DOCX."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data/aurora_land_use_change/final"
FIG = ROOT / "figures/aurora_land_use_change/report"
OUT = ROOT / "reports/aurora_land_use_change"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "16324F"
CYAN = "00A6A6"
GOLD = "F2B134"
LIGHT = "EAF0F4"
MUTED = "64748B"
RED = "D1495B"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", size=font_size)
        shade(table.rows[0].cells[i], NAVY)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            if ridx % 2:
                shade(cells[i], "F5F8FA")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Caption"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_picture(doc, filename, width=6.7, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(FIG / filename), width=Inches(width))
    if caption:
        add_caption(doc, caption)


def add_callout(doc, title, body, color=CYAN):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(title + ". ")
    r.bold = True
    r.font.size = Pt(10)
    r2 = p.add_run(body)
    r2.font.size = Pt(9.5)
    return p


def hyperlink(paragraph, text, url, size=7.7):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(int(size * 2)))
    rpr.append(font_size)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string("263238")
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(31)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor.from_string("000000")
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for level, size in ((1, 19), (2, 13), (3, 10.5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(4)

    for sec in doc.sections:
        header = sec.header.paragraphs[0]
        header.text = "EARTH GENOME  /  AURORA DATA-CENTER LAND-USE CHANGE"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(7.5)
        header.runs[0].font.color.rgb = RGBColor.from_string("000000")
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Evidence current through 8 September 2026   •   ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        footer._p.append(fld)
        for r in footer.runs:
            r.font.size = Pt(7.5)
            r.font.color.rgb = RGBColor.from_string(MUTED)


def page_break(doc):
    doc.add_page_break()


def main():
    sites = pd.read_csv(DATA / "sites.csv").set_index("site_id")
    buildings = pd.read_csv(DATA / "buildings.csv")
    phases = pd.read_csv(DATA / "development_phases.csv")
    evidence = pd.read_csv(DATA / "evidence_catalog.csv")
    dw_built = pd.read_csv(DATA / "dynamic_world_built_2026_summary.csv")
    corroborating_sources = pd.DataFrame(
        [
            ("cmap_2020", "CMAP Land Use Inventory 2020", "https://services5.arcgis.com/LcMXE3TFhi1BSaCY/ArcGIS/rest/services/LUI20_geodatabase_v1_CMAP/FeatureServer/0", "Independent parcel-scale land-use classification", "open_data"),
            ("cmap_2005", "CMAP Land Use Inventory 2005", "https://www.arcgis.com/home/item.html?id=2be439b2e0e84f58ab219b7baf3feff2", "Aerial-photo-derived historical land use at CHI1", "open_data"),
            ("nlcd", "USGS National Land Cover Database 2021", "https://www.usgs.gov/centers/eros/science/annual-national-land-cover-database", "Independent cultivated-crops cross-check", "open_data"),
            ("worldcover", "ESA WorldCover 2021", "https://esa-worldcover.org/en", "Independent 10 m cropland cross-check", "open_data"),
            ("aurora_chi3_plan", "Aurora File 24-0479", "https://aurora-il.legistar.com/LegislationDetail.aspx?GUID=1417FF27-336E-467B-B82E-04ACD7974169&ID=6780393", "Aurora III planning status and final plan", "primary"),
        ],
        columns=evidence.columns,
    )
    evidence = pd.concat([evidence, corroborating_sources], ignore_index=True)

    doc = Document()
    configure_document(doc)

    # Cover.
    p = doc.add_paragraph()
    p.style = doc.styles["Title"]
    p.paragraph_format.space_before = Pt(60)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("LAND-USE CHANGE\nAT AURORA DATA CENTERS")
    r.font.name = "Aptos Display"
    r.font.size = Pt(31)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("000000")
    p2 = doc.add_paragraph("Construction dates, prior land use, disturbed area, current roof footprint, and development timelines")
    p2.runs[0].font.size = Pt(14)
    p2.runs[0].font.color.rgb = RGBColor.from_string("000000")
    p2.paragraph_format.space_after = Pt(18)
    add_picture(doc, "01_study_area_overview.png", width=5.8)
    p = doc.add_paragraph("Study area: northeast Aurora, Illinois  •  41.7795–41.8178° N, 88.2698–88.2092° W")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(8.5)
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    p = doc.add_paragraph("Prepared 8 September 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(NAVY)

    page_break(doc)
    add_heading(doc, "Executive summary", 1)
    doc.add_paragraph(
        "Three distinct data-center campuses were identified inside the supplied GeoJSON: the Edged Chicago campus, the older CyrusOne CHI1-CHI2 campus, and the newer CyrusOne Aurora III/CHI3 campus. Four current roofs have Overture/OSM polygons. The active Edged ORD01-2 roof is represented separately by approved plan dimensions with approximate positioning. The campuses contain approximately 122.7 acres of remotely sensed construction disturbance."
    )
    rows = []
    short_names = {
        "edged_chicago": "Edged Chicago",
        "cyrusone_aurora_chi1_chi2": "CyrusOne CHI1–CHI2",
        "cyrusone_aurora_iii": "CyrusOne Aurora III",
    }
    for sid, row in sites.iterrows():
        rows.append(
            [
                short_names[sid],
                row["first_construction_start"],
                row["predevelopment_land_use"],
                f"{row['disturbed_area_estimate_acres']:.1f} ac\n({row['disturbed_area_lower_acres']:.1f}–{row['disturbed_area_upper_acres']:.1f})",
                (
                    f"{row['mapped_roof_footprint_m2']:,.0f} m² mapped"
                    + (f"\n+ {row['approved_plan_active_roof_m2']:,.0f} m² plan" if row['approved_plan_active_roof_m2'] else "")
                ),
            ]
        )
    add_table(
        doc,
        ["Campus", "Construction began", "Prior land use", "Disturbed area estimate", "Current roof footprint"],
        rows,
        widths=[1.2, 1.0, 1.5, 1.35, 1.35],
        font_size=7.8,
    )
    add_callout(
        doc,
        "The main finding",
        "The two recent campuses occupied undeveloped planned-development parcels that were under active row-crop cultivation immediately before construction. Multiple independent datasets identify approximately 83% to 92% of each parcel as cultivated crops or cropland. The older CHI1-CHI2 campus had a mixed baseline and should not be described as an unqualified agricultural greenfield.",
    )
    add_heading(doc, "Interpretation of area metrics", 2)
    doc.add_paragraph(
        "Roof footprint, gross floor area, campus area, and disturbed area are not interchangeable. Roof footprint is the plan-view building envelope. Gross floor area counts floors. Campus area follows the parcel boundary. Disturbed area includes roofs, pavement, construction pads, utility work, drainage, and exposed ground. Mapped Overture/OSM roofs, approved-plan geometry, and satellite-interpreted built surfaces are therefore maintained as separate provenance classes."
    )
    add_picture(doc, "06_area_comparison.png", width=6.55, caption="Figure 1. Parcel, disturbance, and roof areas. Exact values and uncertainty fields are included in the structured dataset.")

    page_break(doc)
    add_heading(doc, "Study design and evidence", 1)
    doc.add_paragraph(
        "The analysis began with the repository’s harmonized Illinois data-center references and regulatory inventory. Candidate points and polygons were reconciled to city parcel ownership, addresses, current Overture building footprints, municipal development plans, operator announcements, and dated open imagery. This process separated three campuses and resolved the unnamed PNNL polygon as CyrusOne CHI2."
    )
    add_heading(doc, "Imagery and classification workflow", 2)
    methods = [
        ("Dated imagery", "USDA NAIP (1 m; 2011–2023), Landsat Collection 2 (30 m; 2005–2015), and Copernicus Sentinel-2 L2A (10 m; 2016–2026) were inspected as before/during/after evidence."),
        ("Land-cover time series", "Google Earth Engine Dynamic World V1 was summarized annually from 2016 through 2026 using a median probability composite and 10 m argmax class areas."),
        ("Older baseline", "USDA Cropland Data Layer and Landsat were used for CHI1 because Dynamic World begins after that building was constructed."),
        ("Footprints", "Current mapped roofs come from Overture Buildings release 2026-08-19.0, with underlying OSM geometry edits dated October 2025 to February 2026. Edged ORD01-2 is stored separately using approved plan dimensions and approximate positioning."),
        ("Area measurement", "All areas were calculated in NAD83 / UTM zone 16N (EPSG:26916). Spatial interchange files are delivered in EPSG:4326."),
    ]
    add_table(doc, ["Component", "Method"], methods, widths=[1.25, 5.2], font_size=8.3)
    add_heading(doc, "Confidence rules", 2)
    doc.add_paragraph(
        "High confidence requires agreement between dated imagery and a primary documentary source or precise mapped geometry. Medium confidence indicates a clear remote-sensing signal with approximation in timing or geometry. Low-medium confidence is used where a 10–30 m land-cover model is acting as a proxy for a complex, older industrial parcel."
    )
    add_picture(doc, "05_dynamic_world_change.png", width=6.65, caption="Figure 2. Annual Dynamic World class shares show abrupt crop-to-bare-to-built transitions at the recent campuses.")

    page_break(doc)
    add_heading(doc, "Independent checks on prior land use", 1)
    doc.add_paragraph(
        "Dynamic World was checked against four alternative land-cover or land-use sources and direct aerial interpretation. The comparisons use each campus parcel, not a point sample. Agreement is strongest for Edged Chicago and Aurora III, where crop-specific USDA results, regional land-use mapping, and visible field patterns all support recent cultivation."
    )
    add_table(
        doc,
        ["Evidence", "Edged Chicago", "CyrusOne Aurora III"],
        [
            ("Dynamic World", "95.4% crops in 2022", "98.7% crops in 2023"),
            ("USDA CDL", "2022: 42.7% corn, 40.9% soybeans; 84.9% cultivated", "2023: 79.1% corn, 3.6% soybeans; 82.8% cultivated"),
            ("USGS NLCD", "2021: 85.1% cultivated crops", "2021: 84.6% cultivated crops"),
            ("ESA WorldCover", "2021: 83.7% cropland", "2021: 91.9% cropland"),
            ("CMAP inventory", "2020 land use: Agriculture", "2020 land use: Agriculture"),
            ("USDA NAIP", "September 2021: crop rows visible", "August 2023: cultivated field visible"),
            ("Aurora planning record", "January 2023: vacant PDD parcel", "July 2024: vacant PDD parcel"),
        ],
        widths=[1.35, 2.55, 2.55],
        font_size=8,
    )
    add_callout(
        doc,
        "Conclusion for the recent campuses",
        "Confidence is high that both parcels supported active row crops before construction. The defensible range is approximately 84% to 95% agricultural cover at Edged and 83% to 99% at Aurora III. The City records describe the parcels as vacant because that is their development status; it does not rule out interim cultivation. The precise description is therefore undeveloped planned-development land under active row-crop cultivation, rather than agriculturally zoned land.",
    )
    add_heading(doc, "Older CHI1-CHI2 baseline", 2)
    doc.add_paragraph(
        "The older campus requires different wording. CMAP's aerial-photo-derived 2005 inventory assigns the future CHI1 location to vacant residential land use, while the 2006 USDA CDL and Landsat analysis indicates that much of the surface was pasture or cropland. This supports a mixed undeveloped baseline, but not proof of a formal agricultural use. Before CHI2 construction, the September 2015 NAIP image shows wooded cover, open ground, and developed edges. Confidence is medium for the agricultural or pasture share and high for the broader mixed-baseline conclusion."
    )
    doc.add_paragraph(
        "The datasets are not fully independent. USDA CDL uses NLCD information for non-agricultural training, and Dynamic World and WorldCover both use Sentinel imagery. The highest-weight evidence is the agreement between direct NAIP interpretation, CMAP's parcel-scale inventory, and USDA's crop-specific corn and soybean classes."
    )

    page_break(doc)
    add_heading(doc, "Provisional Dynamic World built-surface envelopes", 1)
    doc.add_paragraph(
        "To address map currency, two satellite-interpreted screening layers were extracted directly from the January-August 2026 Dynamic World image collection. Both require built to be the dominant class in a median probability composite. The core layer also requires median built probability of at least 0.50; the inclusive layer uses 0.30. Connected patches smaller than 500 m2 were removed."
    )
    labels = {
        "edged_chicago": "Edged Chicago",
        "cyrusone_aurora_chi1_chi2": "CyrusOne CHI1-CHI2",
        "cyrusone_aurora_iii": "CyrusOne Aurora III",
    }
    dw_rows = []
    for sid in labels:
        core = dw_built[(dw_built.site_id == sid) & (dw_built["product"].eq("core"))].iloc[0]
        inclusive = dw_built[(dw_built.site_id == sid) & (dw_built["product"].eq("inclusive"))].iloc[0]
        dw_rows.append(
            [
                labels[sid],
                f"{core.total_area_acres:.1f} ac / {int(core.polygon_count)} patch(es)",
                f"{inclusive.total_area_acres:.1f} ac / {int(inclusive.polygon_count)} patch(es)",
                str(int(inclusive.large_unmatched_count)),
            ]
        )
    add_table(
        doc,
        ["Campus", "Core built surface", "Inclusive built surface", "Large unmatched candidates"],
        dw_rows,
        widths=[1.75, 1.55, 1.65, 1.4],
        font_size=8,
    )
    add_callout(
        doc,
        "Interpretation",
        "These are provisional built-surface envelopes, not building footprints. At 10 m resolution, Dynamic World commonly merges roofs with pavement, construction pads, and equipment yards. The inclusive layer is useful for flagging potentially missing or recently changed mapped geometry; it should not replace Overture/OSM or an as-built survey.",
    )
    add_picture(doc, "07_dynamic_world_built_envelopes.png", width=6.65, caption="Figure 3. Core and inclusive 2026 Dynamic World built-surface envelopes compared with mapped roofs and approved-plan geometry.")

    page_break(doc)
    add_heading(doc, "Edged Chicago campus", 1)
    e = sites.loc["edged_chicago"]
    doc.add_paragraph(
        "The 65.1-acre Edged campus occupies six Edged Chicago LLC parcels along Bilter and Eola Roads. The approved preliminary plan contains three buildings. ORD01-1 is operating; ORD01-2 was topped out in June 2026 and remains under construction; the third building is planned."
    )
    add_callout(
        doc,
        "Land conversion",
        "Previously undeveloped planned-development land under active row-crop cultivation. Dynamic World assigned 95.4% of the parcel to crops in 2022; USDA CDL identified 83.6% as corn or soybeans and 84.9% as cultivated. CMAP classified the location as Agriculture in 2020, while Aurora's approval record called the parcel vacant. The May 2023 groundbreaking, August 2023 NAIP construction scene, and official February 2025 opening bracket the first phase tightly.",
        GOLD,
    )
    b = buildings[buildings.site_id.eq("edged_chicago")]
    add_table(
        doc,
        ["Building", "Geometry provenance", "Roof footprint", "Reported gross floor area", "Confidence"],
        [[r.building_name, r.geometry_class.replace("_", " "), f"{r.roof_footprint_m2:,.0f} m²", f"{r.reported_gross_floor_area_sqft:,.0f} ft²", r.geometry_confidence] for r in b.itertuples()],
        widths=[1.35, 1.3, 1.2, 1.35, 1.1],
        font_size=8,
    )
    doc.add_paragraph(
        f"Mapped roof area is {e.mapped_roof_footprint_m2:,.0f} m². Adding the active ORD01-2 approved-plan footprint gives a provisional total of {e.current_roof_footprint_m2:,.0f} m². Estimated disturbed area is {e.disturbed_area_estimate_acres:.1f} acres; the plausible upper bound is the full {e.disturbed_area_upper_acres:.1f}-acre parcel assembly."
    )
    add_picture(doc, "02_edged_change.png", width=6.7, caption="Figure 4. Cyan shows the mapped ORD01-1 roof. Purple dashed ORD01-2 geometry uses approved dimensions but approximate positioning.")
    add_heading(doc, "Development timeline", 2)
    add_table(
        doc,
        ["Date", "Event", "Evidence"],
        [
            ("14 Jun 2022", "Last clear preconstruction satellite baseline", "Sentinel-2; crop-dominant parcel"),
            ("22 May 2023", "Campus / ORD01-1 groundbreaking", "Edged official release"),
            ("18 Aug 2023", "Building and broad grading visible", "USDA NAIP"),
            ("27 Feb 2025", "ORD01-1 opened", "Edged official release"),
            ("13 Nov 2025", "ORD01-2 groundbreaking", "Edged official release"),
            ("4 Jun 2026", "ORD01-2 topped out", "Edged official release"),
            ("Q2 2027", "ORD01-2 target operation", "Operator schedule; future"),
        ],
        widths=[1.0, 2.8, 2.6],
        font_size=8,
    )

    page_break(doc)
    add_heading(doc, "CyrusOne CHI1–CHI2 campus", 1)
    c = sites.loc["cyrusone_aurora_chi1_chi2"]
    doc.add_paragraph(
        "The original CyrusOne campus occupies the 41.5-acre 2905 Diehl Road parcel. CHI1 is the former CME Group data center; CHI2 is the smaller western building that appeared after the reported December 2016 construction start. The repository’s previously unnamed PNNL building is therefore reconciled to CHI2."
    )
    add_callout(
        doc,
        "Land conversion",
        "Mixed baseline, not a simple greenfield label. Around the future CHI1 footprint, the 2006 Cropland Data Layer was approximately 74% pasture/cropland, 18% previously developed, 3% forest, and 4% other. The future CHI2 footprint was a wooded/open/developed mix before construction.",
        GOLD,
    )
    b = buildings[buildings.site_id.eq("cyrusone_aurora_chi1_chi2")]
    add_table(
        doc,
        ["Building", "Construction chronology", "Roof footprint", "Reported gross floor area"],
        [
            ("CHI1 / CME", "Change begins 2006–2007; reported year built 2009", f"{b.iloc[0].roof_footprint_m2:,.0f} m²", "428,000 ft²"),
            ("CHI2", "Started Dec 2016; roof visible Jul 2017", f"{b.iloc[1].roof_footprint_m2:,.0f} m²", "316,000 ft²"),
        ],
        widths=[1.25, 2.6, 1.2, 1.3],
        font_size=8,
    )
    doc.add_paragraph(
        f"Estimated disturbed area: {c.disturbed_area_estimate_acres:.1f} acres. Because Dynamic World sees nearly the entire mature industrial parcel as built in 2023, the estimate is best treated as a parcel-scale proxy; the structured dataset carries a 32.0–{c.disturbed_area_upper_acres:.1f}-acre range."
    )
    add_heading(doc, "Development timeline", 2)
    add_table(
        doc,
        ["Date", "Event", "Evidence"],
        [
            ("2006", "Future CHI1 footprint still mostly agricultural/pasture", "Landsat and USDA CDL"),
            ("2006–2007", "Initial CHI1 land-cover transition", "Annual CDL/Landsat interval"),
            ("2009", "CHI1 reported year built", "Facility directory; medium confidence"),
            ("15 Mar 2016", "CME announced sale of 428,000 ft² facility to CyrusOne", "CME Group release"),
            ("Dec 2016", "CHI2 construction reported to begin", "Facility reporting; corroborated by imagery"),
            ("2 Jul 2017", "CHI2 roof visible", "USDA NAIP"),
            ("2018", "CHI2 development documented as 316,000 ft²", "Aurora planning record"),
        ],
        widths=[1.0, 3.1, 2.3],
        font_size=7.8,
    )

    page_break(doc)
    add_heading(doc, "CHI1–CHI2 visual evidence", 1)
    doc.add_paragraph(
        "The image sequence separates the campus’s two main construction waves. CHI1 is absent from the coarse 2006 baseline but fully present by the first available NAIP scene. In 2015 the western CHI2 pad remains a mix of open and wooded cover; by July 2017 its roof is clearly visible."
    )
    add_picture(doc, "03_cyrus_original_change.png", width=6.45, caption="Figure 5. CHI1–CHI2 sequence. The 30 m 2006 baseline supports interval dating but not precise site delineation.")
    add_callout(
        doc,
        "Why the older date remains an interval",
        "The available 30 m annual imagery detects the 2006–2007 land-cover change, while documentary evidence reports a 2009 build year. Without the original grading permit or certificate of occupancy, a single exact construction date would imply false precision.",
        CYAN,
    )

    page_break(doc)
    add_heading(doc, "CyrusOne Aurora III / CHI3 campus", 1)
    a = sites.loc["cyrusone_aurora_iii"]
    doc.add_paragraph(
        "CyrusOne’s second Aurora campus occupies a separate 32.2-acre parcel at 2725 Bilter Road. The announced program contains two buildings totaling 446,000 square feet. One current roof is mapped; no second roof was present by the latest 22 August 2026 Sentinel-2 observation."
    )
    add_callout(
        doc,
        "Land conversion",
        "Previously undeveloped planned-development land under active row-crop cultivation. Dynamic World assigned 98.7% to crops in 2023; USDA CDL identified 82.7% as corn or soybeans and 82.8% as cultivated. CMAP classified the location as Agriculture in 2020, while Aurora's 2024 planning record called the parcel vacant. The August 2023 NAIP image visibly confirms cultivation. Construction began in September 2024.",
        GOLD,
    )
    b = buildings[buildings.site_id.eq("cyrusone_aurora_iii")].iloc[0]
    add_table(
        doc,
        ["Current roof", "Status", "Roof footprint", "Planned campus program"],
        [[b.building_name, "Substantially complete", f"{b.roof_footprint_m2:,.0f} m² ({b.roof_footprint_sqft:,.0f} ft²)", "Two buildings / 446,000 ft² gross"]],
        widths=[1.3, 1.45, 1.7, 2.1],
        font_size=8,
    )
    doc.add_paragraph(
        f"Estimated disturbed area: {a.disturbed_area_estimate_acres:.1f} acres in 2025, when Dynamic World assigned 63.4% to built and 27.2% to bare ground. The parcel boundary provides a {a.disturbed_area_upper_acres:.1f}-acre upper bound."
    )
    add_picture(doc, "04_cyrus_iii_change.png", width=6.7, caption="Figure 6. Aurora III imagery sequence shows the crop-to-bare-to-built progression after the September 2024 groundbreaking.")
    add_heading(doc, "Development timeline", 2)
    add_table(
        doc,
        ["Date", "Event", "Evidence"],
        [
            ("18 Aug 2023", "Cultivated field; no construction", "USDA NAIP"),
            ("13 Sep 2024", "Groundbreaking announced", "CyrusOne official release"),
            ("6 Oct 2024", "Grading and early building works visible", "Sentinel-2"),
            ("26 Oct 2025", "First roof appears substantially enclosed", "Sentinel-2"),
            ("25 Jun 2026", "Substantial completion reported", "CyrusOne campus update; naming caveat"),
            ("22 Aug 2026", "One current roof; second building not visible", "Sentinel-2 and Overture"),
        ],
        widths=[1.05, 3.0, 2.35],
        font_size=8,
    )

    page_break(doc)
    add_heading(doc, "Limitations and appropriate use", 1)
    limitations = [
        ("Disturbed area", "Dynamic World is a 10 m semantic land-cover product, not an engineering grading survey. Built and bare pixels omit revegetated or water-management areas and may overclassify bright industrial surfaces. Use the estimate together with its lower/upper bounds."),
        ("Older construction", "CHI1 predates Dynamic World and the available NAIP sequence. Its start is interval-dated from 30 m Landsat/CDL change and a reported 2009 build year; exact permit and commissioning dates remain unresolved."),
        ("Edged ORD01-2 geometry", "The roof polygon is an analytical approximation based on the approved 417.2 x 612.1 ft plan dimensions. Its area is reliable to plan precision, but its georeferenced position is approximate and is stored separately from mapped roofs."),
        ("Status", "Current means evidence available through 8 September 2026. Construction and operating status can change faster than public map releases."),
        ("Land-use percentages", "Percentages from CDL and Dynamic World are classifier outputs and should be treated as approximate composition, especially at mixed edges and for the 30 m older baseline."),
        ("Land cover versus designation", "Observed cultivation does not establish agricultural zoning, enrollment in a farmland program, or a permanent agricultural use. Aurora described both recent parcels as vacant planned-development land."),
    ]
    add_table(doc, ["Issue", "Implication"], limitations, widths=[1.35, 5.15], font_size=8.2)
    add_heading(doc, "Recommended next validation steps", 2)
    for text in [
        "Obtain Aurora grading permits and certificates of occupancy for exact CHI1 and CHI2 construction/operation dates.",
        "Digitize final as-built civil plans or acquire sub-meter 2026 orthophotography for a survey-quality ORD01-2 roof and disturbance boundary.",
        "Repeat the Dynamic World and Sentinel-2 analysis after the next cloud-free season to track Edged ORD01-2 landscaping and any second Aurora III building.",
        "Use the supplied evidence catalog and per-phase confidence fields when integrating the results into a broader data-center inventory.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    add_heading(doc, "Delivered structured data", 2)
    add_table(
        doc,
        ["File", "Contents"],
        [
            ("aurora_data_center_land_use_change.gpkg", "Canonical sites and provenance-separated geometry layers"),
            ("mapped_building_footprints", "Overture/OSM roof polygons with source edit dates"),
            ("approved_plan_footprints", "Approved dimensions with approximate georeferenced position"),
            ("development_phases (.csv, .parquet)", "Event intervals, status, and confidence"),
            ("predevelopment_land_use (.csv, .parquet)", "Phase-level prior-cover estimates"),
            ("dynamic_world_annual (.csv, .parquet)", "Annual 2016–2026 class probabilities and areas"),
            ("dynamic_world_built_core_2026", "Conservative 10 m satellite built-surface screening polygons"),
            ("dynamic_world_built_inclusive_2026", "More sensitive 10 m screening polygons for map-currency review"),
            ("evidence_catalog / source_crosswalk", "Provenance and repository reconciliation"),
        ],
        widths=[2.6, 3.9],
        font_size=8,
    )

    page_break(doc)
    add_heading(doc, "Sources", 1)
    doc.add_paragraph(
        "Primary and open-data sources are listed below. The machine-readable evidence catalog includes the same records and their analytical use."
    )
    for i, row in enumerate(evidence.itertuples(), 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run(f"[{i}] {row.source_name}. ")
        r.bold = True
        r.font.size = Pt(7.7)
        hyperlink(p, str(row.uri_or_path), str(row.uri_or_path) if str(row.uri_or_path).startswith("http") else (ROOT / str(row.uri_or_path)).as_uri())
        r = p.add_run(f" — {row.use}")
        r.font.size = Pt(7.7)

    add_heading(doc, "Reproducibility", 2)
    doc.add_paragraph(
        "Analysis scripts are in analysis/aurora_land_use_change. Raw imagery chips and manifests are in data/aurora_land_use_change/raw/imagery. Script 08 reproduces the parcel-level USDA CDL, USGS NLCD, and ESA WorldCover cross-checks. The Earth Engine scripts record their datasets, aggregation methods, scale, and project metadata. The final README documents coordinate systems and field semantics."
    )
    add_callout(
        doc,
        "Bottom line",
        "Edged Chicago and CyrusOne Aurora III replaced planned-development parcels that were actively cultivated before construction. CHI1-CHI2 developed in two waves from a mixed vacant, agricultural or pasture, wooded, and partly developed baseline. Disturbed area is approximately 122.7 acres. Mapped roofs total 22.5 acres; including the active Edged plan yields 28.4 acres provisionally.",
        CYAN,
    )

    path = OUT / "aurora_data_center_land_use_change_report.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
